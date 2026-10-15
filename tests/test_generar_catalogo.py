"""Pruebas del catalogo de datos (sin depender de MongoDB real)."""

import sys
from datetime import datetime
from pathlib import Path

from bson import ObjectId

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))

import generar_catalogo as gc  # noqa: E402


class _Cursor:
    def __init__(self, valor):
        self._valor = valor

    def next(self):
        return self._valor


class FakeCollection:
    def __init__(self, docs, index_info=None, count=None, size=None):
        self._docs = docs
        self._index_info = index_info or {"_id_": {"key": [("_id", 1)]}}
        self._count = len(docs) if count is None else count
        self._size = size

    def find(self, filter=None, limit=None):  # noqa: A002
        return iter(self._docs[:limit] if limit else self._docs)

    def index_information(self):
        return self._index_info

    def count_documents(self, filter):  # noqa: A002
        return self._count

    def aggregate(self, pipeline):
        return _Cursor({"storageStats": {"count": self._count, "size": self._size}})


class FakeDB:
    def __init__(self, colecciones):
        self._colecciones = colecciones

    def list_collection_names(self):
        return list(self._colecciones)

    def __getitem__(self, nombre):
        return self._colecciones[nombre]


def _usuario():
    return {
        "_id": ObjectId(),
        "name": "Ana Ruiz",
        "email": "ana@uni.edu.co",
        "birthdate": datetime(2000, 1, 15),
        "location": {"city": "Bogota", "address": {"street": "CL 1 #2-3"}},
        "preferences": {"favoritos": ["sci-fi", "drama"]},
        "movies": [ObjectId(), ObjectId()],
        "imdb": {"rating": 7.5, "votes": 1200},
    }


def _comentario(movie_id):
    return {
        "_id": ObjectId(),
        "movie_id": movie_id,
        "text": "Excelente pelicula",
        "date": datetime(2024, 3, 1),
        "email": "ana@uni.edu.co",
    }


def _db_de_prueba():
    movie_id = ObjectId()
    return FakeDB({
        "users": FakeCollection([_usuario()], count=185, size=20_000),
        "comments": FakeCollection(
            [_comentario(movie_id), _comentario(movie_id)],
            index_info={
                "_id_": {"key": [("_id", 1)]},
                "movie_id_1": {"key": [("movie_id", 1)]},
                "date_-1": {"key": [("date", -1)]},
            },
            count=41079, size=9_000_000,
        ),
        "dim_usuario": FakeCollection([{"_id": "U-1", "nombre": "Ana"}], count=236),
        "fact_interacciones": FakeCollection([{"_id": "C-1", "id_usuario": "U-1"}], count=41079),
        "interacciones_analiticas": FakeCollection([{"_id": "A-1"}], count=41079),
    })


# --------------------------------------------------------------------- puras


def test_aplanar_documento_despliega_anidados():
    campos = dict(gc.aplanar_documento(_usuario()))
    assert "location.address.street" in campos
    assert campos["location.address.street"] == "CL 1 #2-3"
    assert "preferences.favoritos" in campos
    assert "movies" in campos


def test_tipificar_valores_standard():
    assert gc.tipificar_valor(ObjectId()) == "String (UUID)"
    assert gc.tipificar_valor(datetime(2024, 1, 1)) == "Timestamp"
    assert gc.tipificar_valor([1, 2]) == "Array"
    assert gc.tipificar_valor({"a": 1}) == "Objeto"
    assert gc.tipificar_valor(True) == "Boolean"
    assert gc.tipificar_valor(7) == "Integer"
    assert gc.tipificar_valor(7.5) == "Double"
    assert gc.tipificar_valor("x") == "String"
    assert gc.tipificar_valor(None) == "Null / Vacio"


def test_clasificar_arquitectura_por_convencion():
    assert gc.clasificar_arquitectura("dim_usuario").startswith("Data Warehouse")
    assert gc.clasificar_arquitectura("fact_interacciones").startswith("Data Warehouse")
    assert gc.clasificar_arquitectura("interacciones_analiticas").startswith("OLAP")
    assert gc.clasificar_arquitectura("users") == "OLTP (Transaccional)"


def test_pk_y_fk():
    assert gc.es_primary_key("_id") is True
    assert gc.es_primary_key("movie_id") is False
    assert gc.es_foreign_key("movie_id") is True
    assert gc.es_foreign_key("id_usuario") is True
    assert gc.es_foreign_key("_id") is False
    assert gc.es_foreign_key("imdb.rating") is False


def test_esta_indexado_usa_index_information():
    info = {"movie_id_1": {"key": [("movie_id", 1)]}}
    assert gc.esta_indexado("movie_id", info) is True
    assert gc.esta_indexado("text", info) is False
    assert gc.esta_indexado("_id", info) is True


def test_enriquecer_descripcion_usa_hoja_y_fallback():
    assert "Stock Keeping Unit" in gc.enriquecer_descripcion("items.sku", "ventas")
    assert "movie" in gc.enriquecer_descripcion("movie_id", "comments").lower()
    fallback = gc.enriquecer_descripcion("campo_raro", "users")
    assert "campo_raro" in fallback and "users" in fallback


# ---------------------------------------------------------------- escaneo real


def test_construir_resumen_usa_stats_e_indices():
    db = _db_de_prueba()
    resumen = gc.construir_resumen(db, db.list_collection_names())
    usuarios = next(f for f in resumen if f["Coleccion"] == "users")
    comments = next(f for f in resumen if f["Coleccion"] == "comments")
    assert usuarios["Registros"] == 185
    assert usuarios["Rol Arquitectonico"] == "OLTP (Transaccional)"
    assert comments["Indices"] == 3
    assert comments["Tamano Estimado"].endswith("MB")


def test_construir_diccionario_detecciona_claves_indices_y_tipos():
    db = _db_de_prueba()
    filas = gc.construir_diccionario(db, db.list_collection_names())
    por_campo = {(f["Coleccion"], f["Campo"]): f for f in filas}

    assert ("comments", "movie_id") in por_campo
    fk = por_campo[("comments", "movie_id")]
    assert fk["Es Foreign Key (FK)"] == "Si"
    assert fk["Esta Indexado"] == "Si"
    assert fk["Tipo de Dato"] == "String (UUID)"

    anidado = por_campo[("users", "location.address.street")]
    assert anidado["Tipo de Dato"] == "String"
    assert anidado["Es Primary Key (PK)"] == "No"

    pk = por_campo[("users", "_id")]
    assert pk["Es Primary Key (PK)"] == "Si"
    assert pk["Esta Indexado"] == "Si"

    assert por_campo[("users", "imdb.rating")]["Tipo de Dato"] == "Double"
    assert por_campo[("users", "movies")]["Tipo de Dato"] == "Array"
    assert por_campo[("users", "birthdate")]["Tipo de Dato"] == "Timestamp"
    assert por_campo[("fact_interacciones", "id_usuario")]["Es Foreign Key (FK)"] == "Si"


def test_diccionario_ordena_pk_primero_en_cada_coleccion():
    db = _db_de_prueba()
    filas = gc.construir_diccionario(db, ["users"])
    assert filas[0]["Campo"] == "_id"
    campos = [f["Campo"] for f in filas]
    assert campos[1:] == sorted(campos[1:])


# ------------------------------------------------------------------- excel


def test_generar_excel_crea_dos_pestanas(tmp_path):
    from openpyxl import load_workbook

    db = _db_de_prueba()
    resumen = gc.construir_resumen(db, db.list_collection_names())
    diccionario = gc.construir_diccionario(db, db.list_collection_names())
    salida = gc.generar_excel(resumen, diccionario, tmp_path / "catalogo.xlsx")
    assert salida.exists() and salida.stat().st_size > 0

    wb = load_workbook(salida)
    assert wb.sheetnames == ["Resumen Arquitectonico", "Diccionario de Datos"]
    assert [c.value for c in wb["Diccionario de Datos"][1]] == [
        "Coleccion", "Campo", "Tipo de Dato", "Es Primary Key (PK)",
        "Es Foreign Key (FK)", "Esta Indexado", "Descripcion de Negocio",
    ]
    assert wb["Diccionario de Datos"].max_row == len(diccionario) + 1
    assert wb["Resumen Arquitectonico"].freeze_panes == "A2"